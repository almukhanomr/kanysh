"""Export immutable evidence, preserving statuses; never recompute or certify geology."""

import hashlib
import html
import json
import math
from pathlib import Path

from opensubsurface.errors import require


def scalar(value):
    if value is None:
        return "NOT AVAILABLE"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    if isinstance(value, float):
        require(math.isfinite(value), "INVALID_REPORT", "Non-finite report value")
    return value


def load_reports(objects):
    reports = []
    for obj in objects:
        path = Path(obj["paths"]["report.json"])
        require(
            path.stat().st_size <= 32 * 1024**2,
            "REPORT_TOO_LARGE",
            "Report JSON exceeds 32 MiB export budget",
        )
        report = json.loads(path.read_text())
        require(isinstance(report, dict), "INVALID_REPORT", "Report must be a JSON object")
        reports.append((obj, report))
    return reports


def tables(reports):
    result = {}
    for index, (obj, report) in enumerate(reports, 1):
        for name, rows in report.get("tables", {}).items():
            require(
                isinstance(rows, list) and all(isinstance(r, dict) for r in rows),
                "INVALID_REPORT",
                "Expected named row tables",
            )
            target = result.setdefault(name, [])
            for row in rows:
                target.append(dict(report_index=index, report_id=obj["id"], revision=obj["revision"], **row))
        # Older calculation reports have no tables mapping; preserve their full evidence below.
    return result


def workbook(path, reports, tabular, title):
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook(write_only=True)
    all_tables = {
        "Readme": [
            dict(
                title=title,
                meaning="Immutable numerical results snapshot; no live geological formulas",
                missing="NOT AVAILABLE is never zero",
                certification="Numerical checks do not establish geological fitness",
            )
        ],
        "Checks": [
            dict(
                report_index=i,
                id=obj["id"],
                revision=obj["revision"],
                content_hash=obj["content_hash"],
                status=report.get("status", "NOT RUN"),
                validation_status=report.get("validation_status", "NOT RUN"),
            )
            for i, (obj, report) in enumerate(reports, 1)
        ],
        "Assumptions": [
            dict(report_index=i, assumption=value)
            for i, (_, report) in enumerate(reports, 1)
            for value in report.get("assumptions", [])
        ],
        "Definitions": [
            dict(report_index=i, metric=k, definition=v)
            for i, (_, report) in enumerate(reports, 1)
            for k, v in report.get("metric_definitions", {}).items()
        ],
    }
    for name in ("scenarios", "regions"):
        all_tables[name.title()] = tabular.get(name, [])
    for name, rows in tabular.items():
        if name not in ("scenarios", "regions"):
            all_tables[name] = rows
    expected = {}
    used = set()
    for number, (name, rows) in enumerate(all_tables.items(), 1):
        clean = "".join(c if c not in "[]:*?/\\" else "_" for c in name)[:26] or "Table"
        sheet_name = clean if clean.lower() not in used else clean + str(number)
        used.add(sheet_name.lower())
        sheet = book.create_sheet(sheet_name)
        columns = list(dict.fromkeys(k for row in rows for k in row)) or ["status"]
        require(
            len(rows) < 1048576 and len(columns) <= 16384,
            "EXPORT_LIMIT",
            "Excel sheet limits exceeded; split input reports explicitly",
        )
        values = [[scalar(row.get(k)) for k in columns] for row in rows] or [["NO ROWS / NOT RUN"]]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(values) + 1}"
        for col, key in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(col)].width = min(42, max(15, len(key) + 3))
        for n, row in enumerate([columns] + values):
            cells = []
            for value in row:
                require(
                    not isinstance(value, str) or len(value) <= 32767,
                    "EXPORT_LIMIT",
                    "Excel text exceeds cell limit; see source JSON",
                )
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Never execute source text as a formula.
                cell.font = Font(name="Calibri", size=11, color="FFFFFF" if n == 0 else "17344A", bold=n == 0)
                cell.fill = PatternFill(
                    "solid", fgColor="17344A" if n == 0 else "F0F6FA" if n % 2 else "FFFFFF"
                )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(value, float):
                    cell.number_format = '#,##0.000000;[Red](#,##0.000000);"-"'
                cells.append(cell)
            sheet.row_dimensions[n + 1].height = max(
                24,
                min(
                    220,
                    16
                    * max(
                        (len(str(value)) // max(10, min(42, len(columns[i]) + 3)) + 1)
                        for i, value in enumerate(row)
                    ),
                ),
            )
            sheet.append(cells)
        expected[sheet_name] = [columns] + values
    book.save(path)
    # Reopen the actual ZIP via an independent reader pass; compare every cell, not just shape.
    check = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    checked = 0
    try:
        require(check.sheetnames == list(expected), "EXPORT_VALIDATION_FAILED", "Workbook sheet mismatch")
        for name, rows in expected.items():
            actual = list(check[name].iter_rows())
            require(len(actual) == len(rows), "EXPORT_VALIDATION_FAILED", "Workbook row count mismatch")
            for cells, values in zip(actual, rows, strict=True):
                for cell, value in zip(cells, values, strict=True):
                    require(cell.data_type != "f", "EXPORT_VALIDATION_FAILED", "Unexpected Excel formula")
                    if isinstance(value, float):
                        okay = isinstance(cell.value, (float, int)) and math.isclose(
                            cell.value, value, rel_tol=2e-15, abs_tol=0
                        )
                    else:
                        okay = cell.value == value
                    require(okay, "EXPORT_VALIDATION_FAILED", "Workbook value differs from source report")
                    checked += 1
    finally:
        check.close()
    return dict(
        status="PASS",
        checked_cells=checked,
        method="Reopen XLSX and compare every value/type; formula injection forbidden",
        scientific_validation="Inherited from source reports; no new certification",
    )


def pdf(path, reports, tabular, p, slices):
    from matplotlib import font_manager
    from pypdf import PdfReader
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        KeepTogether,
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    font = "OSS-DejaVu"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, font_manager.findfont("DejaVu Sans")))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    cell_style = ParagraphStyle("Cell", fontName=font, fontSize=7, leading=10, wordWrap="CJK")
    story = []

    def paragraph(value, style="BodyText"):
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles[style])

    story.extend(
        [
            paragraph(p["title"], "Title"),
            Spacer(1, 12),
            paragraph(
                "Immutable numerical evidence. Source statuses and NOT RUN checks are preserved. Not recoverable reserves or a production forecast."
            ),
            Spacer(1, 10),
        ]
    )
    for i, (obj, report) in enumerate(reports, 1):
        story.extend(
            [
                paragraph(f"Report {i}: {report.get('operation', obj['id'])}", "Heading2"),
                paragraph(
                    f"Status: {report.get('status', 'NOT RUN')}; independent validation: {report.get('validation_status', 'NOT RUN')}"
                ),
                paragraph(f"{obj['id']} r{obj['revision']} · SHA-256 {obj['content_hash']}"),
            ]
        )
        for label in ("scope", "assumptions", "limitations", "metric_definitions", "independent_checks"):
            if label in report:
                story.append(paragraph(label, "Heading3"))
                value = report[label]
                if isinstance(value, list):
                    for item in value:
                        text = (
                            "; ".join(f"{k}: {scalar(v)}" for k, v in item.items())
                            if isinstance(item, dict)
                            else str(item)
                        )
                        story.append(paragraph(text))
                elif isinstance(value, dict):
                    story.extend(paragraph(f"{k}: {scalar(v)}") for k, v in value.items())
                else:
                    story.append(paragraph(value))
        for name in ("tornado.png", "s_curve.png", "volumes.png"):
            if name in obj["paths"]:
                source = obj["paths"][name]
                w, h = ImageReader(source).getSize()
                width = min(600, w)
                height = width * h / w
                if height > 390:
                    width *= 390 / height
                    height = 390
                story.append(
                    KeepTogether([paragraph(name, "Heading3"), Image(source, width=width, height=height)])
                )
    for index, obj in enumerate(slices, 1):
        import numpy as np
        from matplotlib import pyplot as plt

        section_path = Path(obj["paths"]["slice.npy"])
        require(
            section_path.stat().st_size <= 128 * 1024**2,
            "EXPORT_LIMIT",
            "Slice exceeds 128 MiB export budget",
        )
        section = np.load(section_path, allow_pickle=False)
        require(
            section.ndim == 2 and np.isfinite(section).all(), "INVALID_SLICE", "Expected finite 2D amplitudes"
        )
        fig, ax = plt.subplots(figsize=(9, 4))
        plot = ax.imshow(section.T, aspect="auto", cmap="RdBu_r")
        ax.set(
            xlabel="Section index 0",
            ylabel="Section index 1",
            title=f"{obj['summary']['axis']} {obj['summary']['index']} · {obj['summary']['domain']} · index axes",
        )
        fig.colorbar(plot, ax=ax, label="Amplitude (source units)")
        fig.tight_layout()
        image_path = path.parent / f"slice_{index}.png"
        fig.savefig(image_path, dpi=120)
        plt.close(fig)
        story.append(
            KeepTogether(
                [
                    paragraph(f"Seismic slice {index}", "Heading2"),
                    paragraph(
                        f"{obj['id']} r{obj['revision']} · SHA-256 {obj['content_hash']}. Axis values are array indices; no spatial transform is inferred."
                    ),
                    Image(str(image_path), width=600, height=267),
                ]
            )
        )
    truncations = []
    for name, rows in tabular.items():
        heading = paragraph(name, "Heading2")
        columns = list(dict.fromkeys(k for row in rows for k in row))
        visible = rows[: p["max_pdf_rows_per_table"]]
        if len(visible) < len(rows):
            truncations.append(dict(table=name, shown=len(visible), total=len(rows)))
            story.append(
                paragraph(
                    f"Showing {len(visible)} of {len(rows)} rows (explicit PDF presentation limit). Full tables are in source JSON / XLSX."
                )
            )
        # Split wide tables into column groups; repeat identifiers in every group.
        identifiers = [k for k in ("report_index", "scenario", "region_id") if k in columns]
        rest = [k for k in columns if k not in identifiers and k not in ("report_id", "revision")]
        for start in range(0, len(rest) or 1, 5):
            selected = identifiers + rest[start : start + 5]
            if not selected:
                continue
            content = [[Paragraph(html.escape(k), cell_style) for k in selected]]
            content += [
                [Paragraph(html.escape(str(scalar(row.get(k)))), cell_style) for k in selected]
                for row in visible
            ]
            widths = [740 / len(selected)] * len(selected)
            table = LongTable(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0EDF1")),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD7E0")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ]
                )
            )
            story.extend(
                [
                    KeepTogether([heading, table]) if len(visible) <= 10 else heading,
                    *([] if len(visible) <= 10 else [table]),
                    Spacer(1, 10),
                ]
            )

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawString(40, 20, "Kanysh Basic · immutable results · page " + str(doc.page))

    SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        rightMargin=40,
        leftMargin=40,
        topMargin=35,
        bottomMargin=35,
        title=p["title"],
        author="Kanysh Basic",
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    reader = PdfReader(path)
    text = "\n".join(page.extract_text() for page in reader.pages)
    require(
        len(reader.pages) > 0 and p["title"] in text, "EXPORT_VALIDATION_FAILED", "PDF title or pages missing"
    )
    for obj, report in reports:
        require(
            obj["id"] in text and report.get("validation_status", "NOT RUN") in text,
            "EXPORT_VALIDATION_FAILED",
            "PDF lost provenance/status",
        )
    return dict(
        status="PASS",
        pages=len(reader.pages),
        method="Independent pypdf page/text provenance check",
        presentation_truncations=truncations,
        rendered_layout="NOT RUN; automated tests do not certify layout on all report sizes",
    )


def export(objects, p, out, slices=()):
    reports = load_reports(objects)
    tabular = tables(reports)
    validations, files = {}, {}
    for format in p["formats"]:
        name = "results.xlsx" if format == "xlsx" else "results.pdf"
        path = out / name
        validations[format] = (
            workbook(path, reports, tabular, p["title"])
            if format == "xlsx"
            else pdf(path, reports, tabular, p, slices)
        )
        validations[format]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        files[name] = path
    result = dict(
        operation="results_export",
        status="EXPORTED",
        validation_status="PASS",
        export_checks=validations,
        source_statuses=[
            dict(
                id=obj["id"],
                revision=obj["revision"],
                content_hash=obj["content_hash"],
                status=r.get("status", "NOT RUN"),
                validation_status=r.get("validation_status", "NOT RUN"),
            )
            for obj, r in reports
        ],
        assumptions=p["notes"],
        limitations=[
            "Export PASS verifies file serialization, not source calculation correctness.",
            "No source NOT RUN/FAIL status is promoted to scientific PASS.",
            "PDF row limits are presentation limits; complete numerical tables remain in source JSON and XLSX.",
        ],
        tables={},
    )
    from opensubsurface.adapters.copilot_report import write_report

    files.update(write_report(out, p["title"], result))
    return (
        "analysis_report",
        dict(
            operation="results_export",
            status="EXPORTED",
            validation_status="PASS",
            scientific_validation="See source_statuses in report.json",
        ),
        files,
    )
